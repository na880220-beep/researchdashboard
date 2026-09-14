"""
dv_run.py — orchestrator + Excel writer
=======================================

    python dv_run.py --universe spx
    python dv_run.py --universe all --out ./out
    python dv_run.py --universe r2k --refresh-fundamentals
    python dv_run.py --selftest              # synthetic data, no network

Output is one .xlsx per universe, Google Sheets compatible (values + Arial +
autofilter + frozen panes; no post-2007 array functions).
"""

from __future__ import annotations

import argparse
import glob
import os
import sys
from datetime import datetime

import numpy as np
import pandas as pd

from dv_core import (CFG, Config, fetch_fundamentals_bulk, fetch_prices,
                     load_universe, price_features)
from dv_metrics import (TRIGGER_COLS, apply_triggers, derive_metrics,
                        piotroski, score_frame, stage0_pass, stage1_pass)


# =============================================================================
# PIPELINE
# =============================================================================

def run_universe(uni: str, cfg: Config = CFG, universe_file: str | None = None,
                 tickers: list[str] | None = None,
                 injected: dict | None = None) -> dict:
    print(f"\n=== {uni.upper()} ===")
    if injected is not None:
        tickers = injected["tickers"]
        prices = injected["prices"]
    else:
        tickers = tickers or load_universe(uni, universe_file, cfg)
        print(f"[1/6] universe: {len(tickers)} tickers")
        prices = fetch_prices(tickers, cfg=cfg)

    close, volume = prices["close"], prices["volume"]

    # ---- Stage 0 : bulk price screen ---------------------------------------
    print("[2/6] stage 0 — price / liquidity / neglect")
    px_map, rejects = {}, []
    for t in tickers:
        if t not in close.columns:
            rejects.append({"ticker": t, "stage": 0, "reason": "no_data"})
            continue
        px = price_features(close[t], volume[t])
        ok, why = stage0_pass(px, None, cfg)
        if ok:
            px_map[t] = px
        else:
            rejects.append({"ticker": t, "stage": 0, "reason": why})
    print(f"      {len(px_map)} of {len(tickers)} survive stage 0")

    if not px_map:
        return {"candidates": pd.DataFrame(), "near": pd.DataFrame(),
                "rejects": pd.DataFrame(rejects), "universe_n": len(tickers),
                "stage0_n": 0}

    # ---- Stage 1-3 : fundamentals on survivors only ------------------------
    print(f"[3/6] fundamentals for {len(px_map)} names (cached {cfg.fundamentals_ttl_days}d)")
    if injected is not None:
        fmap = injected["fundamentals"]
    else:
        fmap = fetch_fundamentals_bulk(list(px_map), cfg)

    print("[4/6] stage 1 hard cuts + F-Score")
    rows, near_rows = [], []
    for t, px in px_map.items():
        fd = fmap.get(t) or {}
        if not fd.get("ok"):
            rejects.append({"ticker": t, "stage": 1, "reason": "fundamentals_unavailable"})
            continue

        mcap_ok, why = stage0_pass(px, fd.get("market_cap"), cfg)
        if not mcap_ok:
            rejects.append({"ticker": t, "stage": 0, "reason": why})
            continue

        m = derive_metrics(fd, px)
        f, comps = piotroski(fd)
        m["fscore"] = f
        m.update(comps)
        m["ticker"] = t

        ok, why = stage1_pass(m, cfg)
        if not ok:
            rejects.append({"ticker": t, "stage": 1, "reason": why})
            continue
        if f < cfg.min_fscore:
            m["near_reason"] = f"fscore_{int(f)}"
            near_rows.append(m)
            continue
        rows.append(m)

    print(f"      {len(rows)} candidates, {len(near_rows)} near-miss (F<{cfg.min_fscore})")

    if not rows:
        return {"candidates": pd.DataFrame(), "near": pd.DataFrame(near_rows),
                "rejects": pd.DataFrame(rejects), "universe_n": len(tickers),
                "stage0_n": len(px_map)}

    # ---- Stage 4 : score ----------------------------------------------------
    print("[5/6] scoring")
    cand = score_frame(pd.DataFrame(rows), cfg)

    # ---- Stage 5 : triggers vs last snapshot -------------------------------
    print("[6/6] triggers")
    prev = load_prev_snapshot(uni, cfg)
    cand = apply_triggers(cand, prev, cfg)
    save_snapshot(uni, cand, cfg)

    return {"candidates": cand, "near": pd.DataFrame(near_rows),
            "rejects": pd.DataFrame(rejects), "universe_n": len(tickers),
            "stage0_n": len(px_map)}


# =============================================================================
# SNAPSHOT STORE  (this is what makes "daily" mean something)
# =============================================================================

def save_snapshot(uni: str, df: pd.DataFrame, cfg: Config = CFG) -> None:
    os.makedirs(cfg.snapshot_dir, exist_ok=True)
    keep = [c for c in ("ticker", "composite", "value_score", "turn_score",
                        "fscore", "price", "sector", "norm_earn_yield",
                        "revision_breadth") if c in df.columns]
    path = os.path.join(cfg.snapshot_dir,
                        f"{uni}_{datetime.now():%Y-%m-%d}.csv")
    df[keep].to_csv(path, index=False)


def load_prev_snapshot(uni: str, cfg: Config = CFG) -> pd.DataFrame | None:
    files = sorted(glob.glob(os.path.join(cfg.snapshot_dir, f"{uni}_*.csv")))
    today = os.path.join(cfg.snapshot_dir, f"{uni}_{datetime.now():%Y-%m-%d}.csv")
    files = [f for f in files if f != today]
    if not files:
        return None
    try:
        return pd.read_csv(files[-1])
    except Exception:
        return None


# =============================================================================
# EXCEL OUTPUT
# =============================================================================

DISPLAY = [
    ("ticker", "Ticker", None, 9),
    ("name", "Company", None, 26),
    ("sector", "Sector", None, 20),
    ("industry", "Industry", None, 26),
    ("composite", "Composite", "0.000", 10),
    ("value_score", "Value", "0.000", 9),
    ("turn_score", "Turn", "0.000", 9),
    ("fscore", "F-Score", "0", 8),
    ("sector_rank", "SecRk", "0", 7),
    ("price", "Price", "$#,##0.00", 10),
    ("market_cap", "Mkt Cap ($mm)", "#,##0", 13),
    ("dd_3y_high", "DD 3Y High", "0.0%", 11),
    ("days_below_200dma_1y", "Days <200DMA", "0", 12),
    ("dist_200dma", "vs 200DMA", "0.0%", 10),
    ("norm_earn_yield", "Norm Earn Yld", "0.0%", 13),
    ("norm_pe_proxy", "Norm EV/EBIT", "0.0x", 12),
    ("ev_sales", "EV/Sales", "0.00x", 10),
    ("ev_ebitda", "EV/EBITDA", "0.0x", 10),
    ("fcf_yield", "FCF Yld (3Y)", "0.0%", 11),
    ("p_tangible_bv", "P/TangBV", "0.00x", 10),
    ("net_debt_ebitda", "NetDebt/EBITDA", "0.0x", 13),
    ("interest_coverage", "Int Cov", "0.0x", 9),
    ("current_ratio", "Curr Ratio", "0.00", 10),
    ("share_growth_3y", "Shares 3Y", "0.0%", 10),
    ("rev_cagr_5y", "Rev CAGR 5Y", "0.0%", 11),
    ("rev_yoy_q0", "Rev YoY (Q0)", "0.0%", 11),
    ("gm_q0", "GM (Q0)", "0.0%", 9),
    ("revision_breadth", "Rev Breadth", "0.00", 11),
    ("eps_trend_90d_chg", "EPS Est 90D", "0.0%", 11),
    ("short_pct_float", "Short %Fl", "0.0%", 9),
    ("triggers", "Triggers Today", None, 42),
]


def _write_sheet(wb, title: str, df: pd.DataFrame, cols=None, note: str | None = None):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(title[:31])
    hdr_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    body_font = Font(name="Arial", size=10)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(bottom=thin)

    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(name="Arial", size=9, italic=True, color="595959")
        r0 = 3

    if df is None or df.empty:
        ws.cell(r0, 1, "(no rows)").font = body_font
        return ws

    if cols is None:
        cols = [(c, c, None, 14) for c in df.columns]
    cols = [c for c in cols if c[0] in df.columns]

    for j, (_key, label, _fmt, width) in enumerate(cols, start=1):
        cell = ws.cell(r0, j, label)
        cell.font, cell.fill = hdr_font, hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.row_dimensions[r0].height = 28

    for i, (_, row) in enumerate(df.iterrows(), start=r0 + 1):
        for j, (key, _label, fmt, _w) in enumerate(cols, start=1):
            v = row.get(key)
            if isinstance(v, (np.bool_, bool)):
                v = "Y" if v else ""
            elif isinstance(v, (np.integer,)):
                v = int(v)
            elif isinstance(v, (np.floating, float)):
                v = None if not np.isfinite(v) else float(v)
            if key == "market_cap" and isinstance(v, float):
                v = v / 1e6
            cell = ws.cell(i, j, v)
            cell.font, cell.border = body_font, border
            if fmt and isinstance(v, (int, float)):
                cell.number_format = fmt

    ws.freeze_panes = ws.cell(r0 + 1, 2)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(cols))}{r0 + len(df)}"
    return ws


def sector_rollup(cand: pd.DataFrame) -> pd.DataFrame:
    if cand.empty:
        return pd.DataFrame()
    g = cand.groupby("sector")
    out = pd.DataFrame({
        "sector": g.size().index,
        "n_candidates": g.size().values,
        "median_composite": g["composite"].median().values,
        "median_dd_3y": g["dd_3y_high"].median().values,
        "median_norm_earn_yield": g["norm_earn_yield"].median().values,
        "median_ev_sales": g["ev_sales"].median().values,
        "median_fscore": g["fscore"].median().values,
        "median_rev_yoy": g["rev_yoy_q0"].median().values,
        "pct_margin_improving": g["gm_improving_2q"].mean().values,
        "n_triggered": g["trigger_count"].apply(lambda s: int((s > 0).sum())).values,
        "top_names": [
            ", ".join(cand[cand["sector"] == sec]
                      .nlargest(5, "composite")["ticker"].tolist())
            for sec in g.size().index
        ],
    })
    return out.sort_values("n_candidates", ascending=False)


ROLLUP_COLS = [
    ("sector", "Sector", None, 24),
    ("n_candidates", "Names", "0", 8),
    ("n_triggered", "Triggered", "0", 10),
    ("median_composite", "Med Composite", "0.000", 13),
    ("median_dd_3y", "Med DD 3Y", "0.0%", 11),
    ("median_norm_earn_yield", "Med Norm Yld", "0.0%", 12),
    ("median_ev_sales", "Med EV/Sales", "0.00x", 12),
    ("median_fscore", "Med F", "0.0", 8),
    ("median_rev_yoy", "Med Rev YoY", "0.0%", 12),
    ("pct_margin_improving", "% GM Improving", "0%", 13),
    ("top_names", "Top 5 by Composite", None, 34),
]

HANDOFF_COLS = [
    ("ticker", "Ticker", None, 9),
    ("name", "Company", None, 26),
    ("sector", "Sector", None, 20),
    ("industry", "Industry", None, 28),
    ("composite", "Composite", "0.000", 10),
    ("dd_3y_high", "DD 3Y", "0.0%", 9),
    ("norm_earn_yield", "Norm Earn Yld", "0.0%", 12),
    ("fscore", "F", "0", 6),
    ("net_debt_ebitda", "ND/EBITDA", "0.0x", 10),
    ("rev_cagr_5y", "Rev CAGR 5Y", "0.0%", 11),
    ("rev_yoy_q0", "Rev YoY Q0", "0.0%", 10),
    ("gm_q0", "GM Q0", "0.0%", 8),
    ("triggers", "Triggers", None, 34),
    ("q_cyclical_or_structural", "Cyclical / Structural?", None, 20),
    ("q_what_broke", "What broke (1 line)", None, 34),
    ("q_what_fixes_it", "What fixes it / catalyst", None, 34),
    ("q_verdict", "Verdict", None, 14),
]


def build_workbook(uni: str, res: dict, cfg: Config, path: str) -> str:
    from openpyxl import Workbook

    cand = res["candidates"]
    wb = Workbook()
    wb.remove(wb.active)

    trig = (cand[cand["trigger_count"] > 0]
            .sort_values(["trigger_count", "composite"], ascending=[False, False])
            if not cand.empty else pd.DataFrame())

    _write_sheet(wb, "01_Triggered", trig, DISPLAY,
                 note="TODAY'S ACTIONABLE LIST. Names already in the deep-value universe "
                      "that fired at least one trigger. This is the sheet to read every morning.")

    top = cand.head(cfg.top_n_per_universe) if not cand.empty else cand
    _write_sheet(wb, "02_Candidates", top, DISPLAY,
                 note=f"Full ranked list, top {cfg.top_n_per_universe}. Passed all survival "
                      f"and dilution hard cuts and F-Score >= {cfg.min_fscore}. "
                      "Rebuild monthly; triggers run daily against this.")

    _write_sheet(wb, "03_Sector_Rollup", sector_rollup(cand), ROLLUP_COLS,
                 note="Where the distress is clustered. A sector with many candidates is "
                      "either a cycle trough (buy) or a structural break (avoid) — that "
                      "judgment is qualitative and belongs in sheet 04.")

    handoff = top.copy() if not top.empty else pd.DataFrame()
    for c in ("q_cyclical_or_structural", "q_what_broke", "q_what_fixes_it", "q_verdict"):
        if not handoff.empty:
            handoff[c] = ""
    _write_sheet(wb, "04_Qualitative_Handoff", handoff, HANDOFF_COLS,
                 note="The quantitative screen cannot tell a cyclical trough from a structural "
                      "decline. Fill the last four columns per name (or per industry) before "
                      "committing capital. Blank cells are intentional.")

    near = res["near"]
    if not near.empty:
        near = near.sort_values("fscore", ascending=False)
    _write_sheet(wb, "05_Near_Miss", near.head(80) if not near.empty else near, DISPLAY,
                 note=f"Cheap and survivable but F-Score < {cfg.min_fscore}: fundamentals not "
                      "improving YET. Worth re-checking after each earnings report; a name "
                      "moving from F=4 to F=7 is exactly the setup being hunted.")

    rej = res["rejects"]
    if not rej.empty:
        rej = (rej.groupby(["stage", "reason"]).size()
               .reset_index(name="n").sort_values("n", ascending=False))
    _write_sheet(wb, "06_Reject_Reasons", rej, None,
                 note="Diagnostic. If one reason is eliminating most of the universe, that "
                      "threshold is doing all the work — check it is doing it for the right reason.")

    stats = pd.DataFrame([
        {"metric": "run_timestamp", "value": datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
        {"metric": "universe", "value": uni.upper()},
        {"metric": "universe_size", "value": res["universe_n"]},
        {"metric": "passed_stage0_price", "value": res["stage0_n"]},
        {"metric": "passed_all_cuts", "value": 0 if cand.empty else len(cand)},
        {"metric": "triggered_today", "value": 0 if cand.empty else len(trig)},
        {"metric": "near_miss", "value": len(res["near"])},
    ])
    _write_sheet(wb, "07_Run_Stats", stats, None, note="Run summary.")
    _write_sheet(wb, "08_Parameters", cfg.dump(), None,
                 note="Every threshold used in this run. Change these in dv_core.Config.")

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)
    return path


# =============================================================================
# CLI
# =============================================================================

def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="US deep-value / turnaround screener")
    ap.add_argument("--universe", default="spx", help="ndx | spx | r2k | all")
    ap.add_argument("--universe-file", default=None, help="CSV whose first column is tickers")
    ap.add_argument("--out", default="./out")
    ap.add_argument("--refresh-fundamentals", action="store_true")
    ap.add_argument("--min-fscore", type=int, default=None)
    ap.add_argument("--dd", type=float, default=None, help="e.g. -0.45")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args(argv)

    cfg = Config()
    if a.min_fscore is not None:
        cfg.min_fscore = a.min_fscore
    if a.dd is not None:
        cfg.dd_from_3y_high = a.dd
    if a.refresh_fundamentals:
        cfg.fundamentals_ttl_days = 0

    if a.selftest:
        from dv_selftest import synthetic_injection
        inj = synthetic_injection()
        res = run_universe("test", cfg, injected=inj)
        p = build_workbook("test", res, cfg, os.path.join(a.out, "selftest.xlsx"))
        print(f"\nselftest workbook -> {p}")
        return 0

    unis = ["ndx", "spx", "r2k"] if a.universe == "all" else [a.universe]
    for u in unis:
        res = run_universe(u, cfg, a.universe_file)
        p = build_workbook(u, res, cfg,
                           os.path.join(a.out, f"deepvalue_{u}_{datetime.now():%Y%m%d}.xlsx"))
        print(f"-> {p}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
